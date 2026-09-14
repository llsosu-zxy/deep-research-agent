from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MANUAL_DIR = ROOT / "docs" / "manual"
OUTPUT = ROOT / "docs" / "Deep_Research_Agent_代码级技术说明.docx"

BODY_FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"
TABLE_HEADER_FILL = "1F3864"
TABLE_ALT_FILL = "F2F2F2"
TABLE_BORDER = "D9D9D9"


def set_run_font(
    run,
    name: str = BODY_FONT,
    size: float = 11,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str | None = None,
) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def set_style_font(style, name: str = BODY_FONT) -> None:
    style.font.name = name
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def set_table_borders(table, color: str = TABLE_BORDER) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9, color="666666")
    run.add_text("第 ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run.add_text(" 页")


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    set_style_font(normal)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Title", 24), ("Subtitle", 13), ("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 12.5), ("Heading 4", 11.5)):
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Subtitle"
        set_style_font(style)
        style.paragraph_format.space_before = Pt(10 if name.startswith("Heading") else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_bullet(doc: Document, text: str, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_code_block(doc: Document, lines: list[str], language: str = "") -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    shade_paragraph(paragraph, "F2F2F2")
    run = paragraph.add_run()
    set_run_font(run, name=CODE_FONT, size=8.5)
    for index, line in enumerate(lines):
        if index:
            run.add_break()
        run.add_text(line)


def add_table(doc: Document, rows: list[list[str]], caption: str | None = None) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    rows = [row + [""] * (column_count - len(row)) for row in rows]
    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        set_run_font(r, size=9.5, italic=True, color="595959")
        p.paragraph_format.space_after = Pt(4)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    max_lengths = [
        max(len(str(row[i])) for row in rows)
        for i in range(column_count)
    ]
    total_width = 6.6
    min_width = 0.7
    weights = [max(1.0, length) for length in max_lengths]
    raw_widths = [max(min_width, total_width * weight / sum(weights)) for weight in weights]
    scale = total_width / sum(raw_widths)
    widths = [width * scale for width in raw_widths]
    for row_index, row in enumerate(rows):
        if row_index == 0:
            set_repeat_table_header(table.rows[row_index])
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER_FILL)
                color = "FFFFFF"
                bold = True
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, TABLE_ALT_FILL)
                color = "000000"
                bold = False
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9.5, bold=bold, color=color)


def parse_manual_files() -> list[tuple[str, list[str]]]:
    if not MANUAL_DIR.exists():
        raise FileNotFoundError(f"Manual directory not found: {MANUAL_DIR}")
    files = sorted(MANUAL_DIR.glob("*.md"))
    if not files:
        raise FileNotFoundError(f"No manual Markdown files found in {MANUAL_DIR}")
    return [(path.name, path.read_text(encoding="utf-8").splitlines()) for path in files]


def collect_toc(files: list[tuple[str, list[str]]]) -> list[str]:
    headings: list[str] = []
    for _, lines in files:
        for line in lines:
            if line.startswith("# ") or line.startswith("## "):
                headings.append(line.lstrip("#").strip())
    return headings


def build_docx() -> Path:
    files = parse_manual_files()
    toc = collect_toc(files)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    add_page_number_footer(section)
    configure_styles(doc)

    in_code = False
    code_lines: list[str] = []
    in_table = False
    table_rows: list[list[str]] = []
    pending_caption: str | None = None
    toc_inserted = False

    for _, lines in files:
        for raw_line in lines:
            line = raw_line.rstrip()
            if line.startswith("```"):
                if not in_code:
                    in_code = True
                    code_lines = []
                else:
                    add_code_block(doc, code_lines)
                    in_code = False
                    code_lines = []
                continue
            if in_code:
                code_lines.append(line)
                continue
            if line == ":::table":
                in_table = True
                table_rows = []
                continue
            if line == ":::" and in_table:
                add_table(doc, table_rows, pending_caption)
                pending_caption = None
                in_table = False
                continue
            if in_table:
                if line.strip() == "---":
                    continue
                if line.strip():
                    table_rows.append([part.strip() for part in line.split("|")])
                continue
            if not line.strip():
                continue
            if line.startswith("@TITLE:"):
                paragraph = doc.add_paragraph(style="Title")
                run = paragraph.add_run(line.split(":", 1)[1].strip())
                set_run_font(run, size=24, bold=True, color="000000")
                continue
            if line.startswith("@SUBTITLE:"):
                paragraph = doc.add_paragraph(style="Subtitle")
                run = paragraph.add_run(line.split(":", 1)[1].strip())
                set_run_font(run, size=13, color="000000")
                continue
            if line.startswith("@META:"):
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(line.split(":", 1)[1].strip())
                set_run_font(run, size=10, color="404040")
                paragraph.paragraph_format.space_after = Pt(2)
                continue
            if line == "@PAGEBREAK":
                doc.add_page_break()
                continue
            if line == "@TOC":
                if not toc_inserted:
                    add_body_paragraph(doc, "文档结构总览")
                    for heading in toc:
                        add_bullet(doc, heading)
                    toc_inserted = True
                continue
            if line.startswith("@CAPTION:"):
                pending_caption = line.split(":", 1)[1].strip()
                continue
            if line.startswith("#### "):
                paragraph = doc.add_paragraph(style="Heading 4")
                run = paragraph.add_run(line[5:].strip())
                set_run_font(run, size=11.5, bold=True, color="000000")
                continue
            if line.startswith("### "):
                paragraph = doc.add_paragraph(style="Heading 3")
                run = paragraph.add_run(line[4:].strip())
                set_run_font(run, size=12.5, bold=True, color="000000")
                continue
            if line.startswith("## "):
                paragraph = doc.add_paragraph(style="Heading 2")
                run = paragraph.add_run(line[3:].strip())
                set_run_font(run, size=15, bold=True, color="000000")
                continue
            if line.startswith("# "):
                paragraph = doc.add_paragraph(style="Heading 1")
                run = paragraph.add_run(line[2:].strip())
                set_run_font(run, size=18, bold=True, color="000000")
                continue
            if line.startswith("- "):
                add_bullet(doc, line[2:].strip())
                continue
            numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
            if numbered:
                add_bullet(doc, numbered.group(2), numbered=True)
                continue
            add_body_paragraph(doc, line)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_docx()
    print(output)
